from __future__ import annotations

import argparse
import asyncio
import csv
import sys
from collections.abc import Sequence
from dataclasses import asdict
from datetime import date, datetime, timedelta
from decimal import Decimal
from pathlib import Path
from uuid import UUID, uuid4

from option_opportunities_tools import call_metrics_from_market_price
from option_platform.market_data.base import MarketBar
from option_platform.market_data.tsetmc import TsetmcConfig, TsetmcMarketDataProvider
from option_platform.strategies import (
    ExpiryOpportunityStrategy,
    RecommendationKind,
)

DEFAULT_RISK_FREE_RATE = Decimal("0.30")

OUTPUT_HEADERS = {
    "rank": "رتبه در سررسید",
    "symbol": "نماد اختیار",
    "underlying_symbol": "دارایی پایه",
    "expiry": "تاریخ سررسید میلادی",
    "expiry_jalali": "تاریخ سررسید شمسی",
    "days_to_expiry": "روز تقویمی تا سررسید",
    "business_days_to_expiry": "روز کاری تا سررسید",
    "kind": "نوع پیشنهاد",
    "strike": "قیمت اعمال",
    "underlying_price": "آخرین قیمت دارایی پایه",
    "option_price": "قیمت قابل معامله اختیار",
    "average_traded_value_5d": "میانگین ارزش معاملات ۵ روز کاری گذشته (ریال)",
    "entry_profit_percent": "درصد سود یا زیان لحظه ورود",
    "break_even_price": "قیمت سربه‌سر",
}
OUTPUT_HEADERS.update(
    {
        "implied_volatility_percent": "Implied Volatility (annual %)",
        "delta": "Delta",
        "gamma": "Gamma (per IRR)",
        "theta_per_day": "Theta (IRR per calendar day)",
        "vega_per_percent": "Vega (IRR per 1% IV)",
        "rho_per_percent": "Rho (IRR per 1% rate)",
    }
)


def _to_jalali(value: object) -> str:
    if not hasattr(value, "year"):
        return ""
    year, month, day = value.year, value.month, value.day
    month_days = (0, 31, 59, 90, 120, 151, 181, 212, 243, 273, 304, 334)
    adjusted_year = year + 1 if month > 2 else year
    days = (
        355666
        + 365 * year
        + (adjusted_year + 3) // 4
        - (adjusted_year + 99) // 100
        + (adjusted_year + 399) // 400
        + day
        + month_days[month - 1]
    )
    jalali_year = -1595 + 33 * (days // 12053)
    days %= 12053
    jalali_year += 4 * (days // 1461)
    days %= 1461
    if days > 365:
        jalali_year += (days - 1) // 365
        days = (days - 1) % 365
    if days < 186:
        jalali_month, jalali_day = 1 + days // 31, 1 + days % 31
    else:
        jalali_month, jalali_day = 7 + (days - 186) // 30, 1 + (days - 186) % 30
    return f"{jalali_year:04d}-{jalali_month:02d}-{jalali_day:02d}"


def _business_days_to_expiry(as_of: date, expiry: date) -> int:
    """Count Iran market weekdays after as_of, including the expiry date."""
    if expiry <= as_of:
        return 0
    current = as_of + timedelta(days=1)
    count = 0
    while current <= expiry:
        # Python weekday: Thursday=3 and Friday=4.
        if current.weekday() not in (3, 4):
            count += 1
        current += timedelta(days=1)
    return count


def _average_last_five_values(bars: Sequence[MarketBar], as_of: date) -> Decimal | None:
    completed = [bar.value for bar in bars if bar.trading_date < as_of]
    if not completed:
        return None
    recent = completed[-5:]
    return sum(recent, start=Decimal("0")) / Decimal(len(recent))


async def _load_average_values(
    provider: TsetmcMarketDataProvider,
    instrument_ids_by_symbol: dict[str, UUID],
    as_of: date,
    concurrency: int = 10,
) -> dict[str, Decimal | None]:
    semaphore = asyncio.Semaphore(concurrency)

    async def load(symbol: str, instrument_id: UUID) -> tuple[str, Decimal | None]:
        async with semaphore:
            try:
                bars = await provider.get_daily_bars(instrument_id)
            except Exception:
                return symbol, None
        return symbol, _average_last_five_values(bars, as_of)

    pairs = await asyncio.gather(
        *(load(symbol, instrument_id) for symbol, instrument_id in instrument_ids_by_symbol.items())
    )
    return dict(pairs)


def _export_row(
    row: dict[str, object],
    average_traded_value_5d: Decimal | None = None,
    risk_free_rate: Decimal = DEFAULT_RISK_FREE_RATE,
) -> dict[str, object]:
    spot = row["underlying_price"]
    option_price = row["option_price"]
    strike = row["strike"]
    assert isinstance(spot, Decimal)
    assert isinstance(option_price, Decimal)
    assert isinstance(strike, Decimal)
    expiry = row["expiry"]
    days_to_expiry = row["days_to_expiry"]
    assert isinstance(expiry, date)
    assert isinstance(days_to_expiry, int)
    as_of = expiry - timedelta(days=days_to_expiry)
    enriched = dict(row)
    enriched["expiry_jalali"] = _to_jalali(expiry)
    enriched["business_days_to_expiry"] = _business_days_to_expiry(as_of, expiry)
    enriched["average_traded_value_5d"] = average_traded_value_5d
    enriched["entry_profit_percent"] = (spot - option_price - strike) / spot * Decimal("100")
    pricing_days = max(days_to_expiry - 1, 0)
    metrics = call_metrics_from_market_price(
        option_price,
        spot,
        strike,
        Decimal(pricing_days) / Decimal("365"),
        risk_free_rate,
    )
    enriched["implied_volatility_percent"] = (
        metrics.implied_volatility * Decimal("100") if metrics else None
    )
    for name in ("delta", "gamma", "theta_per_day", "vega_per_percent", "rho_per_percent"):
        enriched[name] = getattr(metrics, name) if metrics else None
    return {key: enriched[key] for key in OUTPUT_HEADERS}


def _rank_rows(rows: list[dict[str, object]]) -> list[dict[str, object]]:
    rows.sort(key=lambda row: (row["expiry"], -row["entry_profit_percent"]))
    previous_expiry: object = None
    rank = 0
    for row in rows:
        rank = rank + 1 if row["expiry"] == previous_expiry else 1
        previous_expiry = row["expiry"]
        row["rank"] = rank
    return rows


def _parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Rank today's TSETMC option opportunities")
    parser.add_argument("--limit", type=int, default=50, help="maximum rows to print")
    parser.add_argument("--output", type=Path, help="optional .xlsx or .csv output path")
    parser.add_argument("--fee", default="0", help="fee in IRR per option")
    parser.add_argument(
        "--risk-free-rate",
        default=str(DEFAULT_RISK_FREE_RATE),
        help="annual continuously-compounded rate as a decimal (default: 0.30)",
    )
    return parser


def _write_xlsx(
    path: Path, rows: list[dict[str, object]], risk_free_rate: Decimal = DEFAULT_RISK_FREE_RATE
) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "پیشنهادهای معاملاتی"
    sheet.sheet_view.rightToLeft = True
    sheet.freeze_panes = "A2"
    sheet.append(list(OUTPUT_HEADERS.values()))
    for row in rows:
        sheet.append([row[key] for key in OUTPUT_HEADERS])

    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.auto_filter.ref = sheet.dimensions
    header_keys = tuple(OUTPUT_HEADERS)
    percent_columns = tuple(
        header_keys.index(key) + 1 for key in ("entry_profit_percent", "implied_volatility_percent")
    )
    numeric_columns = tuple(
        header_keys.index(key) + 1
        for key in (
            "strike",
            "underlying_price",
            "option_price",
            "average_traded_value_5d",
            "break_even_price",
            "delta",
            "gamma",
            "theta_per_day",
            "vega_per_percent",
            "rho_per_percent",
        )
    )
    for column in percent_columns:
        for cell in sheet.iter_cols(min_col=column, max_col=column, min_row=2):
            cell[0].number_format = "0.00"
    for column in numeric_columns:
        for cell in sheet.iter_cols(min_col=column, max_col=column, min_row=2):
            cell[0].number_format = "#,##0.00"
    widths = (16, 18, 15, 20, 18, 22, 19, 18, 15, 23, 24, 34, 25, 18, 22, 15, 15, 28, 25, 25)
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.row_dimensions[1].height = 30

    info = workbook.create_sheet("راهنما")
    info.sheet_view.rightToLeft = True
    info.append(
        ["Risk-free rate", f"{risk_free_rate * Decimal('100')}% continuously compounded annual"]
    )
    info.append(
        ["Black-Scholes assumptions", "European call, no dividend, T = calendar days / 365"]
    )
    info.append(["Greek units", "Theta per calendar day; Vega and Rho per one percentage point"])
    info.append(["زمان تولید", datetime.now().astimezone().isoformat(timespec="seconds")])
    info.append(["مبنای خرید", "بهترین قیمت فروشنده (ask)"])
    info.append(["شاخص ورود", "((قیمت سهم − قیمت اختیار − قیمت اعمال) ÷ قیمت سهم) × ۱۰۰"])
    info.append(["تفسیر", "مثبت: سود فعلی؛ منفی: درصد رشد لازم برای ورود به سود"])
    info.append(["هشدار", "این فایل پیشنهاد تحلیلی است و سفارش واقعی ارسال نمی‌کند."])
    info.column_dimensions["A"].width = 22
    info.column_dimensions["B"].width = 65
    workbook.save(path)


async def main() -> int:
    args = _parser().parse_args()
    risk_free_rate = Decimal(args.risk_free_rate)

    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")

    provider = TsetmcMarketDataProvider(uuid4(), TsetmcConfig())
    try:
        snapshot = await provider.snapshot()
        strategy = ExpiryOpportunityStrategy(
            include_long=True,
            include_short=False,
            fee_per_option=Decimal(args.fee),
        )
        opportunities = [
            item
            for item in strategy.analyze(
                snapshot, provider.instruments, provider.underlying_last_prices
            )
            if item.kind is RecommendationKind.LONG_CALL
        ]
        instrument_ids_by_symbol = {
            instrument.symbol: instrument.instrument_id
            for instrument in provider.instruments.values()
            if instrument.symbol in {item.symbol for item in opportunities}
        }
        average_values = await _load_average_values(
            provider,
            instrument_ids_by_symbol,
            snapshot.provider_timestamp.date(),
        )
    finally:
        await provider.aclose()

    rows = _rank_rows(
        [
            _export_row(asdict(item), average_values.get(item.symbol), risk_free_rate)
            for item in opportunities
        ]
    )
    if args.output and rows:
        if args.output.suffix.lower() == ".xlsx":
            _write_xlsx(args.output, rows, risk_free_rate)
        elif args.output.suffix.lower() == ".csv":
            with args.output.open("w", encoding="utf-8-sig", newline="") as output:
                writer = csv.DictWriter(output, fieldnames=OUTPUT_HEADERS)
                writer.writeheader()
                writer.writerows(rows)
        else:
            raise ValueError("output extension must be .xlsx or .csv")

    columns = (
        "rank",
        "symbol",
        "expiry",
        "expiry_jalali",
        "days_to_expiry",
        "business_days_to_expiry",
        "kind",
        "underlying_price",
        "option_price",
        "average_traded_value_5d",
        "entry_profit_percent",
    )
    print("\t".join(columns))
    for row in rows[: max(args.limit, 0)]:
        print("\t".join("-" if row[key] is None else str(row[key]) for key in columns))
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(asyncio.run(main()))
    except (ValueError, OSError) as exc:
        print(f"error: {exc}", file=sys.stderr)
        raise SystemExit(2) from exc
